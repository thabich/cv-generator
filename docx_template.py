# app/services/generate_docx_template.py
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

# Neues Dokument
doc = Document()

# -------------------------------
# Styles
# -------------------------------
styles = doc.styles

heading1 = styles.add_style('Heading1', 1)
heading1.font.name = 'Arial'
heading1.font.size = Pt(16)
heading1.font.bold = True
heading1.font.color.rgb = RGBColor(58, 83, 155)

heading2 = styles.add_style('Heading2', 1)
heading2.font.name = 'Arial'
heading2.font.size = Pt(14)
heading2.font.bold = True
heading2.font.color.rgb = RGBColor(44, 62, 80)

normal = styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(33, 33, 33)

# -------------------------------
# Header
# -------------------------------
section = doc.sections[0]
header = section.header

header_para = header.add_paragraph()
header_para.text = "{{ config.first_name }} {{ config.last_name }}"
header_para.style = heading1
header_para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

for field in ["email", "phone", "website"]:
    header.add_paragraph(f"{{{{ config.{field} }}}}", style=normal)

# -------------------------------
# Zertifikate Tabelle
# -------------------------------
doc.add_paragraph("{{ _('Certificates') }}", style='Heading2')
cert_table = doc.add_table(rows=1, cols=2)
hdr_cells = cert_table.rows[0].cells
hdr_cells[0].text = "{{ _('Image') }}"
hdr_cells[1].text = "{{ _('Info') }}"

# Beispiel-Zeile für docxtpl
row = cert_table.add_row().cells
row[0].text = "{{ c.cols[0] }}"
row[1].text = "{{ c.cols[1] }}"

# Tabellenränder (optional)
def set_table_border(table, color="D5DBDB", size=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)

    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top','left','bottom','right','insideH','insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), str(size))
        border.set(qn('w:color'), color)
        tblBorders.append(border)

    tblPr.append(tblBorders)

set_table_border(cert_table)

# -------------------------------
# Speichern
# -------------------------------
os.makedirs("templates/export", exist_ok=True)
template_path = os.path.join("templates/export", "template.docx")
doc.save(template_path)
print(f"Template erfolgreich erstellt: {template_path}")
